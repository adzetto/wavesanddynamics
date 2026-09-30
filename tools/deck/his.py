"""His slide, read from his PowerPoint: the words, their order, their look.

    python tools/deck/his.py 9          one slide: every shape top to bottom, with
                                        its place (px on the 1920x1080 slide), fill,
                                        and each paragraph's runs (bold, italic,
                                        size in px, colour); his pictures are saved
                                        to build/r10-deck/his/ to look at; the notes
    python tools/deck/his.py 9 --text   only his strings, one per line, as the
                                        checker compares them
    python tools/deck/his.py --labels   the deck's labels in order

The source is Prob__stat__stoch__est.pptx (73 slides). It is read, never
published. Sizes are px on the 1920 wide slide: his 1 pt is 2 px there.

The deck's labels (30 Sep 2026). His slides keep his numbers, 1 to 73. A
slide of ours is lettered after the slide it follows: 65a is the Kalman loop
between his 65 and 66, and 7a, 7b ... come after his 7. LABELS is the deck in
order; OURS names each slide of ours with the reason it is there, and adding
one is one line there. his(label) is his number, or None for a slide of ours;
stem(label) its files' name (s007, s065a).

strings(label) is what render.py's checker holds a slide to: every paragraph
of every text frame and table cell, in the file's order, less the few
artifacts in SKIP (each with the reason). His page number stays as he typed
it: his 66 reads 66.
"""

import argparse
import functools
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
PPTX = os.environ.get("DECK_PPTX", r"C:\Users\lenovo\Downloads\Transfer\Prob__stat__stoch__est.pptx")
OUT = os.path.join(ROOT, "build", "r10-deck", "his")

HIS_COUNT = 73                      # the slides in his file

# The deck's slides that are not in his file, keyed by label, each with why.
# A label is the number of his slide it follows and a letter: "7a" is the
# first after his 7, "7b" the next. A slide of ours has no words of his to
# check; its words are ours (short, no em dash, no spaced en dash).
OURS = {
    "65a": "the Kalman filter's logic as one loop, asked for on 29 Sep 2026 "
           "(\"kalman filter between 65-66 ... kalman filter logic slide\")",
}

_LABEL = re.compile(r"(\d+)([a-z]*)")


def key(label):
    """A label's place: his number, then the letter (7 < 7a < 7b < 8)."""
    m = _LABEL.fullmatch(str(label).strip().lower())
    if not m:
        raise ValueError(f"not a slide label: {label!r}")
    return int(m.group(1)), m.group(2)


for _l in OURS:
    assert _l == _l.lower() and key(_l)[1] and 1 <= key(_l)[0] <= HIS_COUNT, \
        f"a slide of ours is his number and a letter, after one of his slides: {_l!r}"

# the deck in order, and how many slides it has
LABELS = sorted([str(k) for k in range(1, HIS_COUNT + 1)] + list(OURS), key=key)
COUNT = len(LABELS)


def label(x):
    """The deck's label for x (an int, "7a", " 65A "); ValueError if the deck
    has no such slide."""
    s = str(x).strip().lower()
    key(s)
    s = str(int(re.match(r"\d+", s).group(0))) + s.lstrip("0123456789")
    if s not in LABELS:
        raise ValueError(f"the deck has no slide {x!r}" + (
            " (a slide of ours is added to his.OURS first)" if key(s)[1] else ""))
    return s


def stem(x):
    """The name a slide's files share: s004, s065a."""
    n, letters = key(label(x))
    return f"s{n:03d}{letters}"


def his(x):
    """His slide number for the deck's slide x, or None for a slide of ours."""
    s = label(x)
    return None if s in OURS else int(s)


# Text in his file that is not text on his slide. Each is (his slide, shape
# id, paragraph text, why). Nothing else may be left out.
SKIP = [
    (2, 248, "d", "a stray keystroke inside the GOODNESS FIT chevron, drawn over "
                  "\"whether\" in his export"),
    (2, 534, "the mean, or a single chosen quantile",
     "the same line twice, stacked 16 px apart under \"a design value\" "
     "(shapes 534 and 640); it reads once"),
]


@functools.lru_cache(maxsize=1)
def deck():
    from pptx import Presentation
    p = Presentation(PPTX)
    assert len(p.slides) == HIS_COUNT, f"{PPTX} has {len(p.slides)} slides, not {HIS_COUNT}"
    return p


def _px(e):
    return None if e is None else round(e * 1920 / deck().slide_width)


def _rgb(get):
    try:
        c = get()
        if c is None or c.type is None:
            return None
        try:
            return str(c.rgb)
        except Exception:
            return f"theme:{c.theme_color}"
    except Exception:
        return None


def _fill(sh):
    try:
        f = sh.fill
        if f.type == 1:
            return _rgb(lambda: f.fore_color)
    except Exception:
        pass
    return None


def _line(sh):
    try:
        ln = sh.line
        if ln.fill.type == 1:
            w = ln.width.pt if ln.width else None
            return f"{_rgb(lambda: ln.color)}/{w}pt"
    except Exception:
        pass
    return None


def _paras(tf):
    out = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            f = r.font
            runs.append({"t": r.text, "b": f.bold, "i": f.italic,
                         "px": round(f.size.pt * 2, 1) if f.size else None,
                         "c": _rgb(lambda: f.color)})
        text = "".join(r["t"] for r in runs)
        if text.strip():
            out.append({"text": text, "runs": runs, "level": p.level})
    return out


def shapes(n):
    """Every shape on the deck's slide n (a label; his slide his(n)), groups
    flattened, in the file's order; none on a slide of ours."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    res = []
    h = his(n)
    if h is None:
        return res

    def walk(group, dx=0, dy=0):
        for sh in group:
            d = {"id": sh.shape_id, "name": sh.name, "kind": str(sh.shape_type).split(" ")[0],
                 "x": _px(sh.left), "y": _px(sh.top), "w": _px(sh.width), "h": _px(sh.height),
                 "fill": _fill(sh), "line": _line(sh), "paras": [], "table": None, "picture": False}
            try:
                d["auto"] = str(sh.auto_shape_type).split(" ")[0]
            except Exception:
                pass
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                res.append(d)
                walk(sh.shapes)
                continue
            if sh.has_text_frame:
                d["paras"] = _paras(sh.text_frame)
            if getattr(sh, "has_table", False) and sh.has_table:
                d["table"] = [[_paras(c.text_frame) for c in row.cells] for row in sh.table.rows]
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                d["picture"] = True
                d["image"] = (sh.image.content_type, sh.image.size)
                d["_blob"] = sh.image.blob
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                d["chart"] = [{"cats": list(p.categories),
                               "series": [(s.name, list(s.values)) for s in p.series]}
                              for p in ch.plots]
            res.append(d)

    walk(deck().slides[h - 1].shapes)
    return res


def _skipped(n, sid, text):
    return any(s == n and i == sid and t == text for s, i, t, _ in SKIP)


def strings(n):
    """His paragraphs on the deck's slide n (a label), in the file's order:
    text frames, then each table cell row by row. Whitespace as he typed it,
    and so his page number: his 66 reads 66."""
    out = []
    for d in shapes(n):
        for p in d["paras"]:
            if not _skipped(his(n), d["id"], p["text"]):
                out.append(p["text"])
        for row in d["table"] or ():
            for cell in row:
                out.extend(p["text"] for p in cell)
    return out


def notes(n):
    if his(n) is None:
        return ""
    s = deck().slides[his(n) - 1]
    return s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""


def pictures(n, out=OUT):
    """Save his pictures on slide n; returns [(path, x, y, w, h)]."""
    os.makedirs(out, exist_ok=True)
    got = []
    k = 0
    for d in shapes(n):
        if d.get("picture"):
            k += 1
            ext = d["image"][0].split("/")[-1].replace("jpeg", "jpg")
            p = os.path.join(out, f"{stem(n)}-pic{k}.{ext}")
            with open(p, "wb") as fh:
                fh.write(d["_blob"])
            got.append((p, d["x"], d["y"], d["w"], d["h"]))
    return got


def dump(n):
    n = label(n)
    if his(n) is None:
        print(f"== slide {n}: ours, not in his file ({OURS[n]})")
        return
    print(f"== slide {n}  (his {his(n)} of {len(deck().slides)})")
    for d in sorted(shapes(n), key=lambda d: ((d["y"] or 0) // 8, d["x"] or 0)):
        head = (f"[{d['id']}] {d['kind']}{'/' + d['auto'] if d.get('auto') and d['auto'] != 'RECTANGLE' else ''}"
                f" at ({d['x']},{d['y']}) {d['w']}x{d['h']}")
        if d["fill"]:
            head += f" fill {d['fill']}"
        if d["line"]:
            head += f" line {d['line']}"
        print(head)
        for p in d["paras"]:
            flag = "  (SKIPPED: not on his slide)" if _skipped(his(n), d["id"], p["text"]) else ""
            fmt = " + ".join(
                f"{'B' if r['b'] else ''}{'I' if r['i'] else ''}{r['px'] or ''}px {r['c'] or ''}".strip()
                for r in p["runs"])
            print(f"    {p['text']!r}   [{fmt}]{flag}")
        for i, row in enumerate(d["table"] or ()):
            print(f"    row {i}: " + " || ".join(" / ".join(p["text"] for p in c) for c in row))
        if d.get("chart"):
            print("    chart:", json.dumps(d["chart"])[:400])
    for p, x, y, w, h in pictures(n):
        print(f"picture at ({x},{y}) {w}x{h}: {p}")
    t = notes(n)
    if t:
        print("-- notes (never shown; for the numbers behind a figure):")
        print("   " + t.replace("\n", "\n   "))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("slide", nargs="?", help="the deck's label: his number (9), or ours (65a)")
    ap.add_argument("--text", action="store_true", help="only his strings, one per line")
    ap.add_argument("--labels", action="store_true", help="the deck's labels in order")
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if a.labels or not a.slide:
        print(" ".join(LABELS))
        print(f"{COUNT} slides: his {HIS_COUNT} and ours: " +
              ", ".join(f"{k} ({v})" for k, v in OURS.items()))
    elif a.text:
        for s in strings(a.slide):
            print(s)
    else:
        dump(a.slide)


if __name__ == "__main__":
    main()
